import streamlit as st

st.set_page_config(
    page_title="AI Vakeel",
    page_icon="⚖️",
    layout="centered"
)

st.title("⚖️ AI Vakeel")
st.subheader("Your Indian Legal Assistant")
st.write("भारतीय कानून की जानकारी, आसान भाषा में।")

language = st.selectbox(
    "Choose Language / भाषा चुनें",
    ["English", "हिन्दी"]
)

mode = st.selectbox(
    "What do you need help with?",
    [
        "Ask a Legal Question",
        "Draft a Legal Notice",
        "Draft a Police Complaint",
        "Draft an RTI Application",
        "Understand a Legal Document"
    ]
)

st.info(
    "This is an early prototype. "
    "It does not replace advice from a qualified Indian lawyer."
)

question = st.text_area(
    "Describe your legal problem / अपनी समस्या लिखें",
    height=180,
    placeholder="Explain your problem here..."
)

details = st.text_area(
    "Additional details (optional)",
    height=100
)

if st.button("Generate / तैयार करें"):
    if not question.strip():
        st.warning("Please describe your problem first.")
    else:
        st.subheader("Your Request")
        st.write("**Language:**", language)
        st.write("**Selected service:**", mode)

        if mode == "Draft a Legal Notice":
            st.markdown("### Legal Notice — Draft Template")
            st.write("**To:** [Recipient's name and address]")
            st.write("**Subject:** Legal notice regarding the matter described below.")
            st.write(question)
            if details.strip():
                st.write(details)
            st.write(
                "You should have this draft reviewed by a qualified "
                "lawyer before sending it."
            )

        elif mode == "Draft a Police Complaint":
            st.markdown("### Police Complaint — Draft Template")
            st.write("**To:** The Station House Officer")
            st.write("**Subject:** Complaint regarding the incident described below.")
            st.write("**Respected Sir/Madam,**")
            st.write(question)
            if details.strip():
                st.write(details)
            st.write(
                "I request that the matter be examined and appropriate "
                "action be taken according to law."
            )
            st.write("**Name:** [Your name]")
            st.write("**Date:** [Date]")
            st.write(
                "Please verify the facts and have the complaint reviewed "
                "before submitting it."
            )

        elif mode == "Draft an RTI Application":
            st.markdown("### RTI Application — Draft Template")
            st.write("**To:** The Public Information Officer")
            st.write("**Subject:** Application under the Right to Information Act, 2005")
            st.write("Please provide the following information:")
            st.write(question)
            if details.strip():
                st.write(details)
            st.write("**Applicant's name:** [Your name]")
            st.write("**Address:** [Your postal address]")
            st.write(
                "Check the relevant authority's official procedure and "
                "applicable fee before submitting."
            )

        elif mode == "Understand a Legal Document":
            st.markdown("### Document Review")
            st.write(
                "Please provide the relevant text of the document "
                "to help identify its main points."
            )
            st.write(question)
            st.write(
                "This prototype does not yet perform automated legal "
                "document analysis."
            )

        else:
            st.markdown("### Legal Question")
            st.write(question)
            st.write(
                "This prototype does not yet have an AI legal-answer engine. "
                "Please consult the official resources below or a qualified lawyer."
            )

        st.caption(
            "This is a general information tool, not a substitute "
            "for professional legal advice."
        )

st.divider()
st.subheader("Official Indian Legal Resources")

st.markdown(
    "- [India Code — Central and State legislation](https://www.indiacode.nic.in/)"
)
st.markdown(
    "- [Supreme Court of India](https://www.sci.gov.in/)"
)
st.markdown(
    "- [eCourts Services](https://ecourts.gov.in/)"
)
st.markdown(
    "- [National Legal Services Authority](https://nalsa.gov.in/)"
)
st.markdown(
    "- [RTI Online — Central Government](https://rtionline.gov.in/)"
)

st.caption("AI Vakeel | Prototype | India")
import os
import streamlit as st
from openai import OpenAI

# Optional PDF support
try:
    from pypdf import PdfReader
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False


# --------------------------------------------------
# PAGE SETTINGS
# --------------------------------------------------

st.set_page_config(
    page_title="AI Vakeel | Indian Legal Assistant",
    page_icon="⚖️",
    layout="wide"
)


# --------------------------------------------------
# STYLING
# --------------------------------------------------

st.markdown("""
<style>
.main-title {
    font-size: 42px;
    font-weight: 800;
    text-align: center;
}
.subtitle {
    text-align: center;
    font-size: 18px;
}
.disclaimer {
    background-color: #fff3cd;
    padding: 14px;
    border-radius: 8px;
}
</style>
""", unsafe_allow_html=True)


# --------------------------------------------------
# LANGUAGE SETTINGS
# --------------------------------------------------

st.sidebar.title("⚖️ AI Vakeel")

language = st.sidebar.selectbox(
    "Language / भाषा",
    ["English", "हिन्दी"]
)

if language == "हिन्दी":
    app_subtitle = "आपका भारतीय कानूनी सहायक"
    ask_label = "कानूनी प्रश्न पूछें"
    draft_label = "कानूनी दस्तावेज़ तैयार करें"
    research_label = "कानूनी विषय समझें"
    document_label = "दस्तावेज़ का विश्लेषण करें"
    input_label = "अपनी समस्या या प्रश्न लिखें"
    generate_label = "उत्तर तैयार करें"
    result_label = "आपका परिणाम"
else:
    app_subtitle = "Your Indian Legal Assistant"
    ask_label = "Ask a Legal Question"
    draft_label = "Draft a Legal Document"
    research_label = "Research a Legal Topic"
    document_label = "Analyze a Legal Document"
    input_label = "Describe your legal problem or question"
    generate_label = "Generate Answer"
    result_label = "Your Result"


# --------------------------------------------------
# HEADER
# --------------------------------------------------

st.markdown(
    '<div class="main-title">⚖️ AI VAKEEL</div>',
    unsafe_allow_html=True
)

st.markdown(
    f'<div class="subtitle">{app_subtitle}</div>',
    unsafe_allow_html=True
)

st.write("")

st.warning(
    "Important: AI-generated legal information may be incorrect or "
    "incomplete. Verify laws, dates, citations, and procedures using "
    "official sources. Consult a qualified Indian lawyer for advice."
)


# --------------------------------------------------
# SIDEBAR INFORMATION
# --------------------------------------------------

st.sidebar.markdown("---")
st.sidebar.subheader("Services")

mode = st.sidebar.radio(
    "Choose a service",
    [
        ask_label,
        draft_label,
        research_label,
        document_label
    ]
)

st.sidebar.markdown("---")
st.sidebar.caption(
    "Designed for Indian citizens, advocates, and law students."
)

st.sidebar.caption(
    "Do not enter passwords, bank details, or unnecessary sensitive "
    "personal information."
)


# --------------------------------------------------
# AI CLIENT
# --------------------------------------------------

api_key = (
    st.secrets.get("OPENAI_API_KEY", "")
    if hasattr(st, "secrets")
    else ""
)

api_key = api_key or os.environ.get("OPENAI_API_KEY", "")

client = OpenAI(api_key=api_key) if api_key else None


# --------------------------------------------------
# DOCUMENT TYPES
# --------------------------------------------------

document_types = [
    "Legal Notice",
    "Police Complaint",
    "RTI Application",
    "Affidavit",
    "Consumer Complaint",
    "Application to a Government Office",
    "General Legal Letter"
]


# --------------------------------------------------
# PROMPT BUILDER
# --------------------------------------------------

def build_prompt(mode, question, details, language, document_type):
    language_instruction = (
        "Respond in Hindi using clear, accessible language."
        if language == "हिन्दी"
        else "Respond in clear, accessible English."
    )

    safety_instruction = """
You are AI Vakeel, a legal information and drafting assistant
focused on Indian law.

Rules:
1. Provide general legal information, not a guarantee of legal advice.
2. Never invent statutes, sections, court decisions, citations, or URLs.
3. If uncertain about a legal provision, clearly say so.
4. Distinguish general information from facts supplied by the user.
5. Identify when the user should consult a qualified Indian advocate.
6. Ask for missing facts when they materially affect the answer.
7. Protect privacy. Do not request unnecessary sensitive details.
8. For urgent danger, recommend contacting local emergency services.
9. Do not claim to have searched live legal databases unless that
   capability is actually available.
10. Consider relevant Indian central and state law, jurisdiction,
    and the date of the question.
11. Explain important procedural deadlines where reliably known.
12. Avoid guaranteeing case outcomes.
"""

    if mode == ask_label:
        task = f"""
Answer the user's legal question about Indian law.

Include, where appropriate:
- A short answer
- Relevant legal principles
- Practical next steps
- Documents the user may need
- Possible deadlines or jurisdiction issues
- Official sources to verify the information

Do not invent section numbers or citations.

User question:
{question}

Additional details:
{details}
"""

    elif mode == draft_label:
        task = f"""
Prepare a structured first draft of this Indian legal document:
{document_type}

Use placeholders such as [Applicant Name], [Address], [Date],
and [Case Number] wherever details are missing.

Include:
- Appropriate heading and recipient
- Subject
- Relevant facts
- Clear requests or relief sought
- Date and signature placeholders
- A short checklist of facts and documents to verify

Do not invent facts or claim that a document is legally valid
for every situation.

User's description:
{question}

Additional details:
{details}
"""

    elif mode == research_label:
        task = f"""
Explain this legal topic in the context of Indian law.

Organize the answer into:
1. Overview
2. Main legal principles
3. Important distinctions or exceptions
4. Practical examples
5. Questions that depend on the state or jurisdiction
6. Official sources the user should check

Do not fabricate case law, citations, or current legal amendments.

Topic:
{question}

Additional details:
{details}
"""

    else:
        task = f"""
Help the user understand a legal document or excerpt.

Explain:
- The document's apparent purpose
- Important clauses and terminology
- Rights, duties, or risks that may be relevant
- Deadlines or action items explicitly mentioned
- Questions to ask a qualified advocate

Do not assume the document is authentic or complete.
Do not make a definitive legal conclusion from incomplete text.

Document text or question:
{question}

Additional details:
{details}
"""

    return safety_instruction + "\n" + language_instruction + "\n" + task


# --------------------------------------------------
# DOCUMENT INPUTS
# --------------------------------------------------

document_type = None

if mode == draft_label:
    document_type = st.selectbox(
        "Choose document type",
        document_types
    )


# --------------------------------------------------
# USER INPUT
# --------------------------------------------------

question = st.text_area(
    input_label,
    height=200,
    placeholder=(
        "Example: My employer has not paid my salary for two months."
        if language == "English"
        else "उदाहरण: मेरे नियोक्ता ने दो महीने से वेतन नहीं दिया है।"
    )
)

details = st.text_area(
    "Additional details / अतिरिक्त जानकारी",
    height=100,
    placeholder=(
        "State, dates, relevant facts, and what help you need. "
        "Avoid unnecessary personal information."
    )
)


# --------------------------------------------------
# OPTIONAL PDF UPLOAD
# --------------------------------------------------

uploaded_file = None

if mode == document_label:
    uploaded_file = st.file_uploader(
        "Upload a PDF document (optional)",
        type=["pdf"]
    )

    if uploaded_file and not PDF_SUPPORT:
        st.info(
            "PDF support is not installed yet. You can paste relevant "
            "document text into the question box instead."
        )

    if uploaded_file and PDF_SUPPORT:
        try:
            reader = PdfReader(uploaded_file)
            extracted_pages = []

            for page in reader.pages[:20]:
                extracted_pages.append(page.extract_text() or "")

            extracted_text = "\n".join(extracted_pages)

            if extracted_text.strip():
                st.success("PDF text extracted.")
                question = question + "\n\nPDF TEXT:\n" + extracted_text
            else:
                st.warning(
                    "No selectable text was found. This may be a scanned "
                    "PDF. Paste its text manually if possible."
                )
        except Exception:
            st.error("Could not read this PDF. Please check the file.")


# --------------------------------------------------
# GENERATE ANSWER
# --------------------------------------------------

if st.button(generate_label, type="primary", use_container_width=True):

    if not question.strip():
        st.warning(
            "Please enter your question or describe your problem."
        )

    elif not client:
        st.error(
            "The AI engine is not configured yet. Add your API key "
            "using Streamlit Secrets. Never put the key in GitHub code."
        )

    else:
        prompt = build_prompt(
            mode,
            question,
            details,
            language,
            document_type
        )

        with st.spinner("Preparing your response..."):
            try:
                response = client.responses.create(
                    model="gpt-4.1-mini",
                    input=prompt
                )

                answer = response.output_text

                st.session_state["last_answer"] = answer

            except Exception:
                st.error(
                    "The AI request failed. Check the API key, model "
                    "access, API billing or usage limits, and try again."
                )

        if st.session_state.get("last_answer"):
            st.subheader(result_label)

            st.markdown(st.session_state["last_answer"])

            st.download_button(
                label="Download result as text",
                data=st.session_state["last_answer"],
                file_name="ai_vakeel_result.txt",
                mime="text/plain"
            )


# --------------------------------------------------
# OFFICIAL RESOURCES
# --------------------------------------------------

st.divider()
st.header("Official Indian Legal Resources")

resources = [
    (
        "India Code — Central legislation",
        "https://www.indiacode.nic.in/"
    ),
    (
        "Supreme Court of India",
        "https://www.sci.gov.in/"
    ),
    (
        "eCourts Services",
        "https://ecourts.gov.in/"
    ),
    (
        "National Legal Services Authority",
        "https://nalsa.gov.in/"
    ),
    (
        "RTI Online — Central Government",
        "https://rtionline.gov.in/"
    ),
    (
        "National Consumer Helpline",
        "https://consumerhelpline.gov.in/"
    )
]

for title, url in resources:
    st.markdown(f"- [{title}]({url})")


# --------------------------------------------------
# FOOTER
# --------------------------------------------------

st.divider()

st.caption(
    "AI Vakeel is an experimental legal-information and drafting tool. "
    "It is not a law firm and does not create an advocate-client relationship."
)

st.caption(
    "Always verify current law and consult a qualified lawyer for "
    "important legal decisions."
)
import os
import io
from datetime import date, datetime

import streamlit as st

# Optional integrations
try:
    from openai import OpenAI
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False

try:
    from pypdf import PdfReader
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="AI Vakeel | Indian Legal Assistant",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# PROFESSIONAL INTERFACE
# ============================================================

st.markdown(
    """
    <style>
    .block-container {
        max-width: 1200px;
        padding-top: 1.5rem;
        padding-bottom: 3rem;
    }

    .hero {
        padding: 28px 20px;
        border-radius: 18px;
        text-align: center;
        background: linear-gradient(
            120deg,
            #10233f,
            #184c70
        );
        color: white;
        margin-bottom: 20px;
    }

    .hero h1 {
        font-size: 42px;
        margin-bottom: 8px;
    }

    .hero p {
        font-size: 17px;
        margin-bottom: 0;
    }

    .notice {
        border-left: 5px solid #e0a526;
        padding: 12px;
        background-color: rgba(224, 165, 38, 0.12);
        border-radius: 6px;
        margin: 12px 0;
    }

    .small-muted {
        font-size: 13px;
        opacity: 0.8;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# LANGUAGE
# ============================================================

language = st.sidebar.selectbox(
    "Language / भाषा",
    ["English", "हिन्दी"],
)

HINDI = language == "हिन्दी"

if HINDI:
    subtitle = "आपका भारतीय कानूनी सहायक"
    welcome = "आपकी कानूनी जानकारी, आसान भाषा में।"
    question_label = "अपना प्रश्न या कानूनी समस्या लिखें"
    details_label = "अतिरिक्त जानकारी (वैकल्पिक)"
    generate_label = "उत्तर तैयार करें"
    result_label = "आपका परिणाम"
else:
    subtitle = "Your Indian Legal Assistant"
    welcome = "Legal information and document assistance in simple language."
    question_label = "Describe your legal question or problem"
    details_label = "Additional details (optional)"
    generate_label = "Generate Answer"
    result_label = "Your Result"


# ============================================================
# HEADER
# ============================================================

st.markdown(
    f"""
    <div class="hero">
        <h1>⚖️ AI VAKEEL</h1>
        <p>{subtitle}</p>
        <p>{welcome}</p>
    </div>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="notice">
    <b>Legal disclaimer:</b> AI Vakeel provides general information
    and drafting assistance. It is not a law firm and does not replace
    a qualified Indian advocate. Verify laws, citations, deadlines,
    and procedures before relying on an answer.
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# PRIVATE API KEY CONFIGURATION
# ============================================================

def get_api_key():
    # First try an environment variable.
    key = os.environ.get("OPENAI_API_KEY", "")

    if key:
        return key

    # Then try Streamlit's private Secrets configuration.
    try:
        key = st.secrets.get("OPENAI_API_KEY", "")
        return key
    except Exception:
        return ""


def get_ai_client():
    if not OPENAI_AVAILABLE:
        return None

    api_key = get_api_key()

    if not api_key:
        return None

    try:
        return OpenAI(api_key=api_key)
    except Exception:
        return None


client = get_ai_client()


# ============================================================
# SESSION STATE
# ============================================================

if "case_notes" not in st.session_state:
    st.session_state.case_notes = []

if "last_answer" not in st.session_state:
    st.session_state.last_answer = ""

if "last_prompt" not in st.session_state:
    st.session_state.last_prompt = ""


# ============================================================
# SIDEBAR NAVIGATION
# ============================================================

st.sidebar.title("⚖️ AI Vakeel")
st.sidebar.caption("Indian legal information assistant")

page = st.sidebar.radio(
    "Choose a service",
    [
        "Home",
        "Ask Indian Law",
        "Draft Legal Documents",
        "Legal Research",
        "Case-Law Research",
        "Analyze a PDF",
        "Law Student Mode",
        "Advocate Mode",
        "Case Notes",
        "Important Dates",
        "Official Resources",
    ],
)

st.sidebar.divider()

user_type = st.sidebar.selectbox(
    "User profile",
    [
        "General Citizen",
        "Law Student",
        "Advocate",
    ],
)

st.sidebar.caption(
    "Do not enter passwords, bank details, or unnecessary sensitive "
    "personal information."
)


# ============================================================
# AI RESPONSE FUNCTION
# ============================================================

def ask_ai(prompt):
    """
    Send a prompt to the configured AI model.
    This function does not perform live legal database searches.
    """

    if not OPENAI_AVAILABLE:
        return (
            "The OpenAI Python package is not installed. "
            "Check requirements.txt and redeploy the application."
        )

    if client is None:
        return (
            "The AI service is not configured. Add OPENAI_API_KEY "
            "in Streamlit Secrets. Never put your API key in your "
            "public GitHub code. API usage may incur charges."
        )

    system_instructions = """
You are AI Vakeel, an Indian legal information and drafting assistant.

Your users may be ordinary citizens, law students, and advocates.

IMPORTANT RULES:

1. Focus on Indian law and clearly identify when state-specific law
   or local procedure may matter.

2. Explain legal concepts in accessible language. Follow the language
   requested in the user's prompt.

3. Never invent statutes, sections, judgments, case citations,
   quotations, court orders, or URLs.

4. Do not claim that you searched a live legal database unless an
   actual search tool was used.

5. When you cannot verify whether a law or judgment is current,
   clearly disclose that limitation.

6. Distinguish information supplied by the user from verified facts.

7. Ask for missing facts when they materially affect the answer.

8. For document drafts, use placeholders for unknown names, addresses,
   dates, amounts, and case numbers. Never invent facts.

9. Explain relevant risks, procedural requirements, and possible
   deadlines when reliable information is available.

10. Do not guarantee legal outcomes.

11. Encourage users to consult a qualified Indian advocate when
    the matter is urgent, complex, or has significant consequences.

12. Do not request unnecessary personal information.

13. If the user describes immediate danger, recommend contacting
    appropriate local emergency services.

14. When discussing case law, distinguish verified citations from
    research suggestions that still need verification.

15. Give general legal information, not a definitive lawyer-client
    opinion on incomplete facts.
"""

    try:
        response = client.responses.create(
            model="gpt-4.1-mini",
            instructions=system_instructions,
            input=prompt,
        )

        return response.output_text

    except Exception as exc:
        # Do not expose API keys or sensitive request data.
        return (
            "The AI request failed. Check your private API-key "
            "configuration, model access, usage limits, and network "
            "connection. Technical detail: "
            + type(exc).__name__
        )


# ============================================================
# LANGUAGE INSTRUCTIONS
# ============================================================

def language_instruction():
    if HINDI:
        return (
            "उत्तर सरल और स्पष्ट हिंदी में दें। आवश्यक कानूनी "
            "शब्दों का अर्थ भी समझाएं।"
        )

    return "Respond in clear, accessible English."


# ============================================================
# SHARED QUESTION INPUT
# ============================================================

def question_form(title, prompt_hint, submit_text="Generate Answer"):
    st.subheader(title)

    with st.form(key="form_" + page.replace(" ", "_")):
        question = st.text_area(
            question_label,
            height=190,
            placeholder=prompt_hint,
        )

        details = st.text_area(
            details_label,
            height=100,
        )

        submitted = st.form_submit_button(
            submit_text,
            type="primary",
            use_container_width=True,
        )

    if submitted:
        if not question.strip():
            st.warning("Please enter a question or description.")
            return

        prompt = f"""
Language:
{language_instruction()}

User profile:
{user_type}

User question:
{question}

Additional information:
{details}
"""

        with st.spinner("Preparing your response..."):
            answer = ask_ai(prompt)

        st.session_state.last_answer = answer
        st.session_state.last_prompt = question

        st.subheader(result_label)
        st.markdown(answer)

        show_downloads(answer)


# ============================================================
# DOWNLOAD FUNCTIONS
# ============================================================

def show_downloads(text):
    st.download_button(
        "Download as TXT",
        data=text,
        file_name="ai_vakeel_result.txt",
        mime="text/plain",
        use_container_width=True,
    )

    if DOCX_AVAILABLE:
        try:
            document = Document()
            document.add_heading("AI Vakeel - Legal Draft", level=1)

            for paragraph in text.splitlines():
                document.add_paragraph(paragraph)

            buffer = io.BytesIO()
            document.save(buffer)
            buffer.seek(0)

            st.download_button(
                "Download as Word (.docx)",
                data=buffer.getvalue(),
                file_name="ai_vakeel_result.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True,
            )
        except Exception:
            st.info("Word export is temporarily unavailable.")


# ============================================================
# HOME
# ============================================================

if page == "Home":
    st.header("Welcome to AI Vakeel")

    st.write(
        "AI Vakeel helps people understand legal concepts, organize "
        "legal information, and prepare first drafts of documents."
    )

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("⚖️ Legal Help")
        st.write(
            "- Ask questions about Indian law\n"
            "- Understand legal procedures\n"
            "- Prepare first drafts of legal documents\n"
            "- Organize facts and research questions"
        )

    with col2:
        st.subheader("📚 Learning & Research")
        st.write(
            "- Explore legal topics\n"
            "- Prepare law-study notes\n"
            "- Develop case-law research queries\n"
            "- Summarize supplied legal text"
        )

    st.info(
        "Choose a service from the sidebar to get started. "
        "AI features require a configured API key."
    )

    st.subheader("Official websites")

    st.markdown(
        "- [India Code](https://www.indiacode.nic.in/)\n"
        "- [Supreme Court of India](https://www.sci.gov.in/)\n"
        "- [eCourts Services](https://ecourts.gov.in/)\n"
        "- [National Legal Services Authority](https://nalsa.gov.in/)"
    )


# ============================================================
# ASK INDIAN LAW
# ============================================================

elif page == "Ask Indian Law":
    question_form(
        "Ask a Question About Indian Law",
        (
            "Example: My employer has not paid my salary for two months. "
            "What steps can I consider?"
        ),
    )


# ============================================================
# DOCUMENT DRAFTING
# ============================================================

elif page == "Draft Legal Documents":
    st.header("📝 Legal Document Drafting")

    document_type = st.selectbox(
        "Select document type",
        [
            "Legal Notice",
            "Police Complaint",
            "RTI Application",
            "Affidavit",
            "Consumer Complaint",
            "Application to a Government Office",
            "Salary or Employment Complaint",
            "Land or Property Representation",
            "Rental Dispute Letter",
            "General Legal Letter",
        ],
    )

    st.caption(
        "Drafts require factual review. Some documents have special "
        "formatting, service, filing, or notarization requirements."
    )

    with st.form("draft_document_form"):
        facts = st.text_area(
            "Describe the facts",
            height=200,
            placeholder=(
                "Include the relevant events, dates, amounts, "
                "what happened, and what you want to request."
            ),
        )

        recipient = st.text_input(
            "Recipient or authority (optional)"
        )

        state = st.text_input(
            "Indian state or jurisdiction (optional)"
        )

        desired_outcome = st.text_area(
            "What outcome are you seeking?",
            height=90,
        )

        submitted = st.form_submit_button(
            "Draft Document",
            type="primary",
            use_container_width=True,
        )

    if submitted:
        if not facts.strip():
            st.warning("Please describe the facts first.")
        else:
            prompt = f"""
{language_instruction()}

Prepare a first draft of the following document under the
relevant Indian legal context, where applicable.

Document type:
{document_type}

Recipient:
{recipient or "[Recipient to be confirmed]"}

State or jurisdiction:
{state or "[Jurisdiction to be confirmed]"}

Facts supplied by the user:
{facts}

Desired outcome:
{desired_outcome or "[Desired outcome to be confirmed]"}

Requirements:
- Use a suitable heading and subject.
- Organize the facts clearly and chronologically.
- Use placeholders for missing details.
- Do not invent facts or legal provisions.
- Avoid claiming the document is legally sufficient in every case.
- Include a checklist of items the user must verify.
- Identify whether a lawyer should review it before use.
- Use the selected language.
"""

            with st.spinner("Drafting document..."):
                answer = ask_ai(prompt)

            st.session_state.last_answer = answer
            st.subheader("Your Draft")
            st.markdown(answer)
            show_downloads(answer)


# ============================================================
# LEGAL RESEARCH
# ============================================================

elif page == "Legal Research":
    question_form(
        "📚 Legal Topic Research",
        (
            "Example: Explain the general procedure for filing a "
            "consumer complaint in India."
        ),
    )

    st.caption(
        "This is AI-assisted research, not a live legal database. "
        "Verify current law and official procedures independently."
    )


# ============================================================
# CASE-LAW RESEARCH
# ============================================================

elif page == "Case-Law Research":
    st.header("⚖️ Case-Law Research Assistant")

    st.warning(
        "This feature generates research guidance. It does not "
        "automatically search or verify live court databases."
    )

    with st.form("case_law_form"):
        issue = st.text_area(
            "Legal issue or question",
            height=150,
            placeholder=(
                "Example: Supreme Court decisions discussing "
                "principles of natural justice in administrative action."
            ),
        )

        court = st.selectbox(
            "Preferred court",
            [
                "Any relevant Indian court",
                "Supreme Court of India",
                "Relevant High Court",
                "District or subordinate court",
            ],
        )

        period = st.text_input(
            "Relevant time period (optional)"
        )

        submit = st.form_submit_button(
            "Prepare Research Plan",
            type="primary",
        )

    if submit:
        if not issue.strip():
            st.warning("Enter a legal issue first.")
        else:
            prompt = f"""
{language_instruction()}

Prepare a careful Indian case-law research plan.

Legal issue:
{issue}

Preferred court:
{court}

Time period:
{period or "Not specified"}

Include:
1. Main legal questions.
2. Relevant legal concepts and likely search terms.
3. Possible statutory provisions to investigate, without guessing
   section numbers if uncertain.
4. Official databases or court websites to search.
5. A checklist for verifying judgments, citations, dates, and
   whether a judgment remains good law.
6. A table template with columns for case name, citation, court,
   date, issue, holding, and verification status.

Do not invent case names, citations, quotations, or holdings.
Clearly label unverified research suggestions.
"""

            with st.spinner("Preparing research plan..."):
                answer = ask_ai(prompt)

            st.session_state.last_answer = answer
            st.markdown(answer)
            show_downloads(answer)

    st.markdown("#### Official research starting points")

    st.markdown(
        "- [Supreme Court of India](https://www.sci.gov.in/)\n"
        "- [India Code](https://www.indiacode.nic.in/)\n"
        "- [eCourts Services](https://ecourts.gov.in/)"
    )


# ============================================================
# PDF ANALYSIS
# ============================================================

elif page == "Analyze a PDF":
    st.header("📄 Legal Document Analyzer")

    st.write(
        "Upload a text-based PDF and ask for a plain-language summary. "
        "Scanned documents may not contain extractable text."
    )

    if not PDF_AVAILABLE:
        st.error(
            "PDF support is unavailable. Ensure pypdf is listed in "
            "requirements.txt and redeploy."
        )

    uploaded = st.file_uploader(
        "Upload PDF",
        type=["pdf"],
    )

    question = st.text_area(
        "What would you like to understand?",
        value=(
            "Summarize this document, identify important dates and "
            "obligations, explain difficult terms, and list questions "
            "I should ask a lawyer."
        ),
        height=120,
    )

    if st.button("Analyze Document", type="primary"):
        if not PDF_AVAILABLE:
            st.error("PDF support is not installed.")
        elif uploaded is None:
            st.warning("Please upload a PDF.")
        else:
            try:
                reader = PdfReader(uploaded)

                # Limit extraction to reduce accidental overload.
                page_text = []

                for page_number, page in enumerate(reader.pages[:25]):
                    page_text.append(
                        f"Page {page_number + 1}:\n"
                        + (page.extract_text() or "")
                    )

                extracted = "\n\n".join(page_text)

                if not extracted.strip():
                    st.warning(
                        "No readable text was extracted. The PDF may "
                        "be scanned or image-based."
                    )
                else:
                    # Keep prompt size limited.
                    extracted = extracted[:30000]

                    prompt = f"""
{language_instruction()}

Analyze the following user-supplied legal document text.

User's question:
{question}

Document text:
{extracted}

Provide:
1. A plain-language summary.
2. Important dates, amounts, and obligations actually present.
3. Clauses or passages that may require attention.
4. Unclear or missing information.
5. Questions the user may wish to ask a qualified advocate.

Do not assume authenticity or completeness.
Do not invent text that is absent.
Do not give a definitive legal opinion based on incomplete context.
Identify that the document text may be incomplete.
"""

                    with st.spinner("Analyzing PDF text..."):
                        answer = ask_ai(prompt)

                    st.session_state.last_answer = answer
                    st.markdown(answer)
                    show_downloads(answer)

            except Exception:
                st.error(
                    "The PDF could not be processed. Check whether "
                    "it is a valid PDF and try again."
                )


# ============================================================
# LAW STUDENT MODE
# ============================================================

elif page == "Law Student Mode":
    st.header("🎓 Law Student Study Assistant")

    task = st.selectbox(
        "Study activity",
        [
            "Explain a legal concept",
            "Prepare revision notes",
            "Create practice questions",
            "Compare two legal concepts",
            "Create a case brief from supplied text",
            "Prepare an exam answer outline",
        ],
    )

    with st.form("student_form"):
        topic = st.text_area(
            "Topic or supplied material",
            height=180,
        )

        level = st.selectbox(
            "Study level",
            [
                "Beginner",
                "Undergraduate law student",
                "Advanced study",
            ],
        )

        submit = st.form_submit_button(
            "Create Study Material",
            type="primary",
        )

    if submit:
        if not topic.strip():
            st.warning("Enter a topic or text.")
        else:
            prompt = f"""
{language_instruction()}

You are helping a law student study Indian law.

Task:
{task}

Study level:
{level}

Topic or supplied text:
{topic}

Give a well-organized educational response.
Use examples when helpful.
Distinguish general principles from exceptions.
Do not invent statutory provisions or case citations.
For supplied case text, do not add facts that are not present.
Encourage verification against current course materials and law.
"""

            with st.spinner("Preparing study material..."):
                answer = ask_ai(prompt)

            st.session_state.last_answer = answer
            st.markdown(answer)
            show_downloads(answer)


# ============================================================
# ADVOCATE MODE
# ============================================================

elif page == "Advocate Mode":
    st.header("👨‍⚖️ Advocate Research & Case Preparation")

    task = st.selectbox(
        "Choose a task",
        [
            "Organize case facts",
            "Identify potential legal issues",
            "Prepare a chronology",
            "Generate research questions",
            "Create an argument outline",
            "Prepare client interview questions",
            "Summarize supplied judgments",
        ],
    )

    st.warning(
        "Review all output independently. Do not upload privileged "
        "or confidential material unless you are authorized to do so "
        "and have assessed the service's privacy requirements."
    )

    with st.form("advocate_form"):
        case_material = st.text_area(
            "Facts or legal material",
            height=220,
        )

        jurisdiction = st.text_input(
            "Jurisdiction or court (optional)"
        )

        submit = st.form_submit_button(
            "Prepare Analysis",
            type="primary",
        )

    if submit:
        if not case_material.strip():
            st.warning("Enter facts or legal material first.")
        else:
            prompt = f"""
{language_instruction()}

Assist with organizing legal work under Indian law.

Task:
{task}

Jurisdiction:
{jurisdiction or "Not specified"}

Supplied material:
{case_material}

Requirements:
- Separate supplied facts from assumptions.
- Identify missing facts and evidence.
- Highlight potential counterarguments where relevant.
- Identify research questions requiring verification.
- Do not invent authorities, evidence, or facts.
- Do not claim that a legal proposition is settled without support.
- Label the output as preliminary work product for independent review.
"""

            with st.spinner("Preparing preliminary analysis..."):
                answer = ask_ai(prompt)

            st.session_state.last_answer = answer
            st.markdown(answer)
            show_downloads(answer)


# ============================================================
# CASE NOTES
# ============================================================

elif page == "Case Notes":
    st.header("📋 Case Notes Organizer")

    st.info(
        "These notes are stored only in the current app session. "
        "They are not a secure permanent case-management system."
    )

    with st.form("case_notes_form"):
        case_title = st.text_input(
            "Case title or private reference"
        )

        case_summary = st.text_area(
            "Summary or notes",
            height=160,
        )

        next_action = st.text_input(
            "Next action (optional)"
        )

        submit = st.form_submit_button("Add Note")

    if submit:
        if not case_title.strip() or not case_summary.strip():
            st.warning("Enter a title and notes.")
        else:
            st.session_state.case_notes.append(
                {
                    "title": case_title,
                    "summary": case_summary,
                    "action": next_action,
                    "created": datetime.now().strftime(
                        "%Y-%m-%d %H:%M"
                    ),
                }
            )

            st.success("Note added to this session.")

    st.subheader("Saved in current session")

    if not st.session_state.case_notes:
        st.write("No notes added yet.")
    else:
        for index, note in enumerate(
            st.session_state.case_notes
        ):
            with st.expander(note["title"]):
                st.write("Created:", note["created"])
                st.write(note["summary"])
                st.write("Next action:", note["action"] or "None")

                if st.button(
                    "Delete this note",
                    key=f"delete_note_{index}",
                ):
                    st.session_state.case_notes.pop(index)
                    st.rerun()


# ============================================================
# IMPORTANT DATE TRACKER
# ============================================================

elif page == "Important Dates":
    st.header("📅 Important Date Tracker")

    st.warning(
        "This is a general reminder tool, not a legal limitation "
        "calculator. Legal deadlines may depend on statutes, court "
        "rules, service dates, exclusions, and case-specific facts."
    )

    title = st.text_input(
        "Reminder title"
    )

    event_date = st.date_input(
        "Important date",
        value=date.today(),
    )

    if st.button("Calculate Days", type="primary"):
        days = (event_date - date.today()).days

        if days > 0:
            st.success(
                f"{days} day(s) until the selected date."
            )
        elif days == 0:
            st.info("The selected date is today.")
        else:
            st.warning(
                f"The selected date was {-days} day(s) ago."
            )

        if title.strip():
            st.write("Reminder:", title)

        st.caption(
            "Record this date in a reliable calendar and verify "
            "any legal deadline with a qualified advocate."
        )


# ============================================================
# OFFICIAL RESOURCES
# ============================================================

elif page == "Official Resources":
    st.header("🇮🇳 Official Indian Legal Resources")

    resources = [
        (
            "India Code",
            "Central legislation and statutory materials",
            "https://www.indiacode.nic.in/",
        ),
        (
            "Supreme Court of India",
            "Official Supreme Court website",
            "https://www.sci.gov.in/",
        ),
        (
            "eCourts Services",
            "Court services and case-related information",
            "https://ecourts.gov.in/",
        ),
        (
            "National Legal Services Authority",
            "Legal services and legal-aid information",
            "https://nalsa.gov.in/",
        ),
        (
            "RTI Online",
            "Central Government RTI applications",
            "https://rtionline.gov.in/",
        ),
        (
            "National Consumer Helpline",
            "Consumer grievance information",
            "https://consumerhelpline.gov.in/",
        ),
    ]

    for title, description, url in resources:
        with st.container(border=True):
            st.subheader(title)
            st.write(description)
            st.markdown(f"[Open official website]({url})")


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "AI Vakeel | Indian legal information and drafting prototype"
)

st.caption(
    "AI-generated content may be inaccurate or outdated. "
    "Verify legal claims against authoritative sources."
)

st.caption(
    "Protect confidential information. Use this tool as an aid, "
    "not as a substitute for qualified legal advice."
)

